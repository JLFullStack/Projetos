import subprocess
import subprocess
from pathlib import Path
from collections import defaultdict

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Reference

# ==========================================
# CONFIGURAÇÃO
# ==========================================

REPOSITORIO = Path(
    r"C:\CAMINHO\PARA\O\REPOSITORIO\GIT")

ARQUIVO_SAIDA = "relatorio_de_produtividade_cadastro_base_local.xlsx"

AUTORES_IGNORADOS = {
    "DESENVOLVEDORES PARTICIPANTES DO PROJETO A SEREM IGNORADOS PELO RELATÓRIO",
}

# ==========================================


def executar_git(repo, comando):
    resultado = subprocess.run(
        comando,
        cwd=repo,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="ignore",
    )

    if resultado.returncode != 0:
        raise RuntimeError(resultado.stderr)

    return resultado.stdout


if not (REPOSITORIO / ".git").exists():
    raise Exception(f"Repositório não encontrado: {REPOSITORIO}")

print(f"Processando: {REPOSITORIO.name}")

dados = defaultdict(
    lambda: {
        "commits": 0,
        "adicoes": 0,
        "remocoes": 0,
        "arquivos": set(),
    }
)

log = executar_git(
    REPOSITORIO,
    [
        "git",
        "log",
        "--all",
        "--numstat",
        "--format=AUTOR|%an",
    ],
)

autor_atual = None

for linha in log.splitlines():

    if linha.startswith("AUTOR|"):

        autor = linha.removeprefix("AUTOR|").strip()

        if autor in AUTORES_IGNORADOS:
            autor_atual = None
            continue

        autor_atual = autor

        dados[autor]["commits"] += 1
        continue

    if autor_atual is None:
        continue

    partes = linha.split("\t")

    if len(partes) != 3:
        continue

    try:
        add = int(partes[0])
        rem = int(partes[1])
        arquivo = partes[2]
    except ValueError:
        continue

    dados[autor_atual]["adicoes"] += add
    dados[autor_atual]["remocoes"] += rem
    dados[autor_atual]["arquivos"].add(arquivo)

linhas = [
    {
        "Desenvolvedor": autor,
        "Commits": info["commits"],
        "Linhas Adicionadas": info["adicoes"],
        "Linhas Removidas": info["remocoes"],
        "Linhas Liquidas": info["adicoes"] - info["remocoes"],
        "Arquivos Alterados": len(info["arquivos"]),
    }
    for autor, info in dados.items()
]

df = pd.DataFrame(linhas)

if df.empty:
    print("Nenhum dado encontrado.")
    raise SystemExit()

df.sort_values(by="Linhas Adicionadas", ascending=False, inplace=True)

with pd.ExcelWriter(ARQUIVO_SAIDA, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Resumo", index=False)

wb = load_workbook(ARQUIVO_SAIDA)

ws = wb["Resumo"]

# ==========================================
# GRÁFICO COMMITS
# ==========================================

chart_commits = BarChart()
chart_commits.title = "Commits por Desenvolvedor"

chart_commits.add_data(
    Reference(ws, min_col=2, min_row=1, max_row=len(df) + 1), titles_from_data=True
)

chart_commits.set_categories(
    Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
)

ws.add_chart(chart_commits, "H2")

# ==========================================
# GRÁFICO PARTICIPAÇÃO
# ==========================================

chart_participacao = PieChart()
chart_participacao.title = "Participação por Linhas Adicionadas"

chart_participacao.add_data(
    Reference(ws, min_col=3, min_row=1, max_row=len(df) + 1), titles_from_data=True
)

chart_participacao.set_categories(
    Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
)

ws.add_chart(chart_participacao, "H18")

wb.save(ARQUIVO_SAIDA)

print(f"Relatório gerado: {ARQUIVO_SAIDA}")
